from datetime import date

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.init_db import seed_standard_data
from app.models import Account, Category, Organization, Transaction, TransactionItem
from app.models.enums import Scope, TransactionType
from app.services.accounting_service import AccountingService
from app.services.ledger_service import LedgerService
from app.services.report_service import ReportService


def _seed(db_session: Session) -> tuple[Organization, dict[str, Account], dict[str, Category]]:
    organization = seed_standard_data(db_session)
    db_session.flush()
    accounts = {account.code: account for account in db_session.scalars(select(Account))}
    categories = {category.name: category for category in db_session.scalars(select(Category))}
    return organization, accounts, categories


def _post_expense(
    db_session: Session,
    organization: Organization,
    *,
    amount: int,
    scope: Scope,
    payment_id: int,
    items: list[TransactionItem],
    memo: str,
) -> Transaction:
    transaction = Transaction(
        organization_id=organization.id,
        occurred_on=date(2026, 9, 1),
        transaction_type=TransactionType.EXPENSE,
        scope=scope,
        amount=amount,
        memo=memo,
        payment_account_id=payment_id,
        items=items,
    )
    AccountingService(db_session).post_transaction(transaction)
    return transaction


def test_trial_balance_and_ledgers_use_journal_lines(db_session: Session) -> None:
    organization, accounts, categories = _seed(db_session)
    _post_expense(
        db_session,
        organization,
        amount=50_000,
        scope=Scope.BUSINESS,
        payment_id=accounts["2100"].id,
        items=[
            TransactionItem(
                category_id=categories["사무용품"].id,
                amount=50_000,
                scope=Scope.BUSINESS,
                line_no=1,
            )
        ],
        memo="프린터 용지",
    )
    ledger = LedgerService(db_session)
    trial = ledger.trial_balance(organization_id=organization.id)
    assert trial.is_balanced is True
    assert trial.debit_total == trial.credit_total == 50_000

    book = ledger.transaction_ledger(organization_id=organization.id)
    assert {(row.account_code, row.debit_amount, row.credit_amount) for row in book} == {
        ("5100", 50_000, 0),
        ("2100", 0, 50_000),
    }

    supplies = ledger.account_ledgers(organization_id=organization.id, account_id=accounts["5100"].id)[0]
    assert supplies.lines[-1].balance == 50_000
    assert supplies.closing_balance == 50_000


def test_business_filter_excludes_personal_and_keeps_balance(db_session: Session) -> None:
    organization, accounts, categories = _seed(db_session)
    _post_expense(
        db_session,
        organization,
        amount=50_000,
        scope=Scope.BUSINESS,
        payment_id=accounts["2100"].id,
        items=[
            TransactionItem(category_id=categories["사무용품"].id, amount=50_000, scope=Scope.BUSINESS, line_no=1)
        ],
        memo="사업 사무용품",
    )
    _post_expense(
        db_session,
        organization,
        amount=100_000,
        scope=Scope.MIXED,
        payment_id=accounts["2100"].id,
        items=[
            TransactionItem(account_id=accounts["5100"].id, amount=30_000, scope=Scope.BUSINESS, line_no=1),
            TransactionItem(account_id=accounts["5700"].id, amount=70_000, scope=Scope.PERSONAL, line_no=2),
        ],
        memo="마트 혼합 결제",
    )
    _post_expense(
        db_session,
        organization,
        amount=20_000,
        scope=Scope.PERSONAL,
        payment_id=accounts["1100"].id,
        items=[
            TransactionItem(category_id=categories["식비"].id, amount=20_000, scope=Scope.PERSONAL, line_no=1)
        ],
        memo="개인 식비",
    )

    ledger = LedgerService(db_session)
    all_trial = ledger.trial_balance(organization_id=organization.id, business_only=False)
    business_trial = ledger.trial_balance(organization_id=organization.id, business_only=True)

    assert all_trial.is_balanced is True
    assert all_trial.debit_total == 170_000
    assert business_trial.is_balanced is True
    assert business_trial.debit_total == business_trial.credit_total == 80_000
    personal_use = [row for row in business_trial.rows if row.account_code == "5700"]
    assert personal_use == []
    cash = [row for row in business_trial.rows if row.account_code == "1100"]
    assert cash == []


def test_business_excel_contains_expected_sheets_and_filters_personal(db_session: Session) -> None:
    organization, accounts, categories = _seed(db_session)
    _post_expense(
        db_session,
        organization,
        amount=50_000,
        scope=Scope.BUSINESS,
        payment_id=accounts["2100"].id,
        items=[
            TransactionItem(category_id=categories["사무용품"].id, amount=50_000, scope=Scope.BUSINESS, line_no=1)
        ],
        memo="사업 사무용품",
    )
    _post_expense(
        db_session,
        organization,
        amount=100_000,
        scope=Scope.MIXED,
        payment_id=accounts["2100"].id,
        items=[
            TransactionItem(account_id=accounts["5100"].id, amount=30_000, scope=Scope.BUSINESS, line_no=1),
            TransactionItem(account_id=accounts["5700"].id, amount=70_000, scope=Scope.PERSONAL, line_no=2),
        ],
        memo="마트 혼합 결제",
    )
    _post_expense(
        db_session,
        organization,
        amount=20_000,
        scope=Scope.PERSONAL,
        payment_id=accounts["1100"].id,
        items=[
            TransactionItem(category_id=categories["식비"].id, amount=20_000, scope=Scope.PERSONAL, line_no=1)
        ],
        memo="개인 식비",
    )

    path = ReportService(db_session).generate_business_excel(
        organization_id=organization.id,
        date_from=date(2026, 1, 1),
        date_to=date(2026, 9, 30),
    )
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["요약", "거래내역", "분개장", "계정별원장", "시산표"]

    summary = workbook["요약"]
    labels = {row[0]: row[1] for row in summary.iter_rows(min_row=3, max_col=2, values_only=True) if row[0]}
    assert labels["거래 건수"] == 2
    assert labels["사업 지출"] == 80_000
    assert labels["시산표 균형"] == "일치"

    tx_sheet = workbook["거래내역"]
    memos = [row[5] for row in tx_sheet.iter_rows(min_row=2, values_only=True)]
    amounts = [row[3] for row in tx_sheet.iter_rows(min_row=2, values_only=True)]
    assert "개인 식비" not in memos
    assert "마트 혼합 결제" in memos
    assert 70_000 not in amounts
    assert 30_000 in amounts
    assert 50_000 in amounts

    trial_sheet = workbook["시산표"]
    account_codes = [row[0] for row in trial_sheet.iter_rows(min_row=2, values_only=True) if row[0] not in ("합계", None)]
    assert "5700" not in account_codes
    assert "5100" in account_codes
