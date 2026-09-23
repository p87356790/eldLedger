from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.config import settings
from app.models import Transaction
from app.models.enums import RecordStatus, Scope, TransactionType
from app.schemas.reports import (
    AccountLedger,
    BusinessReportSummary,
    JournalBookEntry,
    TrialBalanceResponse,
)
from app.services.ledger_service import LedgerService

HEADER_FILL = PatternFill("solid", fgColor="1565C0")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE"),
)
NUMBER_FORMAT = "#,##0"
TYPE_LABELS = {
    TransactionType.INCOME: "수입",
    TransactionType.EXPENSE: "지출",
    TransactionType.TRANSFER: "이체",
}
SCOPE_LABELS = {
    Scope.PERSONAL: "개인",
    Scope.BUSINESS: "사업",
    Scope.MIXED: "혼합",
}


class ReportService:
    """세무사 제출용 사업자 Excel 보고서를 Journal 데이터에서 생성한다."""

    def __init__(self, session: Session) -> None:
        self._session = session
        self._ledger = LedgerService(session)

    def build_business_summary(
        self,
        *,
        organization_id: int,
        date_from: date | None = None,
        date_to: date | None = None,
    ) -> BusinessReportSummary:
        income_total, expense_total, transaction_count = self._ledger.business_transaction_amounts(
            organization_id=organization_id,
            date_from=date_from,
            date_to=date_to,
        )
        trial = self._ledger.trial_balance(
            organization_id=organization_id,
            date_from=date_from,
            date_to=date_to,
            business_only=True,
        )
        return BusinessReportSummary(
            date_from=date_from,
            date_to=date_to,
            transaction_count=transaction_count,
            income_total=income_total,
            expense_total=expense_total,
            net_income=income_total - expense_total,
            trial_balance_debit_total=trial.debit_total,
            trial_balance_credit_total=trial.credit_total,
            is_balanced=trial.is_balanced,
        )

    def generate_business_excel(
        self,
        *,
        organization_id: int,
        date_from: date | None = None,
        date_to: date | None = None,
    ) -> Path:
        summary = self.build_business_summary(
            organization_id=organization_id,
            date_from=date_from,
            date_to=date_to,
        )
        transactions = self._business_transactions(
            organization_id=organization_id,
            date_from=date_from,
            date_to=date_to,
        )
        journal = self._ledger.journal_book(
            organization_id=organization_id,
            date_from=date_from,
            date_to=date_to,
            business_only=True,
        )
        ledgers = self._ledger.account_ledgers(
            organization_id=organization_id,
            date_from=date_from,
            date_to=date_to,
            business_only=True,
        )
        trial = self._ledger.trial_balance(
            organization_id=organization_id,
            date_from=date_from,
            date_to=date_to,
            business_only=True,
        )

        workbook = Workbook()
        self._write_summary(workbook.active, summary)
        self._write_transactions(workbook.create_sheet("거래내역"), transactions)
        self._write_journal(workbook.create_sheet("분개장"), journal)
        self._write_account_ledgers(workbook.create_sheet("계정별원장"), ledgers)
        self._write_trial_balance(workbook.create_sheet("시산표"), trial)

        reports_dir = Path(settings.data_dir) / "reports"
        reports_dir.mkdir(parents=True, exist_ok=True)
        period = _period_stamp(date_from, date_to)
        path = reports_dir / f"{period}_사업자_회계자료.xlsx"
        workbook.save(path)
        return path

    def _business_transactions(
        self,
        *,
        organization_id: int,
        date_from: date | None,
        date_to: date | None,
    ) -> list[Transaction]:
        stmt = (
            select(Transaction)
            .options(selectinload(Transaction.items))
            .where(Transaction.organization_id == organization_id)
            .where(Transaction.status != RecordStatus.REVERSED)
            .where(Transaction.scope.in_((Scope.BUSINESS, Scope.MIXED)))
            .order_by(Transaction.occurred_on, Transaction.id)
        )
        if date_from is not None:
            stmt = stmt.where(Transaction.occurred_on >= date_from)
        if date_to is not None:
            stmt = stmt.where(Transaction.occurred_on <= date_to)
        rows: list[Transaction] = []
        for transaction in self._session.scalars(stmt):
            if self._ledger.reported_business_amount(transaction) > 0:
                rows.append(transaction)
        return rows

    def _write_summary(self, sheet: Worksheet, summary: BusinessReportSummary) -> None:
        sheet.title = "요약"
        sheet["A1"] = "eldLedger 사업자 회계자료"
        sheet["A1"].font = Font(size=16, bold=True, color="1565C0")
        period = "전체"
        if summary.date_from or summary.date_to:
            start = summary.date_from.isoformat() if summary.date_from else ""
            end = summary.date_to.isoformat() if summary.date_to else ""
            period = f"{start} ~ {end}"
        rows = [
            ("기간", period),
            ("구분", "사업"),
            ("거래 건수", summary.transaction_count),
            ("사업 수입", summary.income_total),
            ("사업 지출", summary.expense_total),
            ("사업 손익", summary.net_income),
            ("시산표 차변 합계", summary.trial_balance_debit_total),
            ("시산표 대변 합계", summary.trial_balance_credit_total),
            ("시산표 균형", "일치" if summary.is_balanced else "불일치"),
        ]
        for index, (label, value) in enumerate(rows, start=3):
            sheet.cell(index, 1, label).font = Font(bold=True)
            cell = sheet.cell(index, 2, value)
            if isinstance(value, int) and label != "거래 건수":
                cell.number_format = NUMBER_FORMAT
        _set_widths(sheet, [22, 28])

    def _write_transactions(self, sheet: Worksheet, transactions: list[Transaction]) -> None:
        headers = ["날짜", "유형", "구분", "금액", "결제계정ID", "사용처", "메모"]
        _write_header(sheet, headers)
        for row_index, transaction in enumerate(transactions, start=2):
            amount = self._ledger.reported_business_amount(transaction)
            values: list[object] = [
                transaction.occurred_on,
                TYPE_LABELS[transaction.transaction_type],
                SCOPE_LABELS[transaction.scope],
                amount,
                transaction.payment_account_id,
                transaction.merchant or "",
                transaction.memo or "",
            ]
            _write_row(sheet, row_index, values, money_cols={4})
        _set_widths(sheet, [14, 10, 10, 16, 14, 24, 40])

    def _write_journal(self, sheet: Worksheet, entries: list[JournalBookEntry]) -> None:
        headers = ["날짜", "분개ID", "적요", "계정코드", "계정", "차변", "대변"]
        _write_header(sheet, headers)
        row_index = 2
        for entry in entries:
            for line in entry.lines:
                _write_row(
                    sheet,
                    row_index,
                    [
                        entry.occurred_on,
                        entry.journal_entry_id,
                        entry.description or "",
                        line.account_code,
                        line.account_name,
                        line.debit_amount,
                        line.credit_amount,
                    ],
                    money_cols={6, 7},
                )
                row_index += 1
        _set_widths(sheet, [14, 12, 28, 12, 18, 16, 16])

    def _write_account_ledgers(self, sheet: Worksheet, ledgers: list[AccountLedger]) -> None:
        headers = ["계정코드", "계정", "날짜", "적요", "차변", "대변", "잔액"]
        _write_header(sheet, headers)
        row_index = 2
        for ledger in ledgers:
            for line in ledger.lines:
                _write_row(
                    sheet,
                    row_index,
                    [
                        ledger.account_code,
                        ledger.account_name,
                        line.occurred_on,
                        line.description or "",
                        line.debit_amount,
                        line.credit_amount,
                        line.balance,
                    ],
                    money_cols={5, 6, 7},
                )
                row_index += 1
        _set_widths(sheet, [12, 18, 14, 28, 16, 16, 16])

    def _write_trial_balance(self, sheet: Worksheet, trial: TrialBalanceResponse) -> None:
        headers = ["계정코드", "계정", "계정유형", "차변", "대변"]
        _write_header(sheet, headers)
        row_index = 2
        for row in trial.rows:
            _write_row(
                sheet,
                row_index,
                [row.account_code, row.account_name, row.account_type.value, row.debit_total, row.credit_total],
                money_cols={4, 5},
            )
            row_index += 1
        sheet.cell(row_index, 1, "합계").font = Font(bold=True)
        debit_cell = sheet.cell(row_index, 4, trial.debit_total)
        credit_cell = sheet.cell(row_index, 5, trial.credit_total)
        debit_cell.number_format = NUMBER_FORMAT
        credit_cell.number_format = NUMBER_FORMAT
        debit_cell.font = Font(bold=True)
        credit_cell.font = Font(bold=True)
        sheet.cell(row_index + 2, 1, "차변 합계 = 대변 합계")
        sheet.cell(row_index + 2, 2, "일치" if trial.is_balanced else "불일치")
        _set_widths(sheet, [12, 18, 14, 16, 16])


def _write_header(sheet: Worksheet, headers: list[str]) -> None:
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(1, col, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        cell.border = THIN
    sheet.freeze_panes = "A2"


def _write_row(sheet: Worksheet, row_index: int, values: list[object], money_cols: set[int]) -> None:
    for col, value in enumerate(values, start=1):
        cell = sheet.cell(row_index, col, value)
        cell.border = THIN
        if col in money_cols and isinstance(value, int):
            cell.number_format = NUMBER_FORMAT


def _set_widths(sheet: Worksheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _period_stamp(date_from: date | None, date_to: date | None) -> str:
    if date_from and date_to:
        return f"{date_from:%Y%m%d}_{date_to:%Y%m%d}"
    if date_from:
        return f"{date_from:%Y%m%d}_"
    if date_to:
        return f"_{date_to:%Y%m%d}"
    return date.today().strftime("%Y")
